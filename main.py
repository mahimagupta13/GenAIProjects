from __future__ import annotations

import re
from collections import Counter
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

import streamlit as st
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


MAX_FILE_SIZE = 10 * 1024 * 1024
STOP_WORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "for", "from", "in",
    "is", "it", "of", "on", "or", "our", "that", "the", "their", "this",
    "to", "we", "will", "with", "you", "your",
}
GENERIC_KEYWORDS = {
    "ability", "candidate", "company", "experience", "job", "least", "looking",
    "must", "preferred", "required", "requirements", "role", "strong", "team",
    "work", "year", "years",
}
SECTION_PATTERNS = {
    "Contact information": r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b",
    "Professional summary": r"\b(summary|profile|objective)\b",
    "Work experience": r"\b(experience|employment|work history)\b",
    "Education": r"\b(education|academic|university|college)\b",
    "Skills": r"\b(skills|technologies|technical proficiencies)\b",
    "Projects": r"\b(projects|portfolio)\b",
}
ACTION_VERBS = {
    "achieved", "automated", "built", "created", "delivered", "designed",
    "developed", "improved", "increased", "launched", "led", "managed",
    "optimized", "reduced", "streamlined",
}
WORD_PATTERN = re.compile(r"[a-z][a-z0-9+#.-]*", re.IGNORECASE)


def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract readable text from a supported resume format."""
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if suffix == ".docx":
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(paragraphs + table_cells)

    if suffix == ".txt":
        return content.decode("utf-8-sig")

    raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")


def extract_keywords(job_description: str, limit: int = 20) -> list[str]:
    """Find useful single- and multi-word terms in a job description."""
    words = [word.lower().strip(".-") for word in WORD_PATTERN.findall(job_description)]
    counts: Counter[str] = Counter()

    for size in (1, 2, 3):
        for start in range(len(words) - size + 1):
            phrase_words = words[start : start + size]
            if any(word in STOP_WORDS for word in phrase_words):
                continue
            phrase = " ".join(phrase_words)
            if len(phrase) < 2 or (size == 1 and phrase in GENERIC_KEYWORDS):
                continue
            counts[phrase] += 1

    ranked = sorted(
        counts,
        key=lambda phrase: (
            -(counts[phrase] * (1 + 0.25 * (phrase.count(" ") + 1))),
            -len(phrase.split()),
            phrase,
        ),
    )
    return ranked[:limit]


def contains_keyword(text: str, keyword: str) -> bool:
    pattern = rf"(?<![a-z0-9]){re.escape(keyword)}(?![a-z0-9])"
    return re.search(pattern, text, re.IGNORECASE) is not None


def analyze_resume(resume_text: str, job_description: str = "") -> dict[str, object]:
    """Return transparent resume quality checks and optional keyword matching."""
    normalized_text = resume_text.lower()
    sections = {
        name: bool(re.search(pattern, normalized_text, re.IGNORECASE))
        for name, pattern in SECTION_PATTERNS.items()
    }
    action_verb_count = sum(
        len(re.findall(rf"\b{re.escape(verb)}\b", normalized_text))
        for verb in ACTION_VERBS
    )
    has_metrics = bool(re.search(r"\b\d+(\.\d+)?\s*%|\$\s*\d|\b\d+\+?\s*(users|clients|people|projects|hours|days)\b", normalized_text))
    completeness_score = round(
        (sum(sections.values()) / len(sections)) * 70
        + min(action_verb_count, 3) / 3 * 15
        + (15 if has_metrics else 0)
    )

    result: dict[str, object] = {
        "sections": sections,
        "action_verb_count": action_verb_count,
        "has_metrics": has_metrics,
        "completeness_score": completeness_score,
    }

    if job_description.strip():
        keywords = extract_keywords(job_description)
        matched = [
            keyword for keyword in keywords
            if contains_keyword(normalized_text, keyword)
        ]
        missing = [
            keyword for keyword in keywords
            if not contains_keyword(normalized_text, keyword)
        ]
        coverage = len(matched) / len(keywords) if keywords else 0.0
        result.update(
            {
                "keywords": keywords,
                "matched_keywords": matched,
                "missing_keywords": missing,
                "match_score": round(coverage * 70 + completeness_score * 0.3),
            }
        )

    return result


def build_recommendations(analysis: dict[str, object]) -> list[str]:
    recommendations: list[str] = []
    sections = analysis["sections"]
    assert isinstance(sections, dict)

    missing_sections = [name for name, present in sections.items() if not present]
    if missing_sections:
        recommendations.append(
            "Add or clearly label these sections: " + ", ".join(missing_sections) + "."
        )
    if analysis["action_verb_count"] == 0:
        recommendations.append(
            "Start experience bullets with specific action verbs such as "
            "developed, improved, or delivered."
        )
    if not analysis["has_metrics"]:
        recommendations.append(
            "Add measurable outcomes where possible, such as percentages, time saved, "
            "or the number of users supported."
        )
    missing_keywords = analysis.get("missing_keywords", [])
    if missing_keywords:
        recommendations.append(
            "If accurate, include relevant missing job-description terms: "
            + ", ".join(missing_keywords[:8])
            + "."
        )
    if not recommendations:
        recommendations.append(
            "Good coverage. Tailor the wording to each role and keep every claim accurate."
        )
    return recommendations


def main() -> None:
    st.set_page_config(page_title="Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 Resume Checker")
    st.write(
        "Get a quick, practical review of your resume and see how its wording "
        "aligns with a job description."
    )
    st.caption(
        "Your resume is analyzed by this app using transparent text checks; "
        "it is not sent to an external AI service. The scores are guidance, not "
        "a hiring decision or a guarantee of ATS performance."
    )

    uploaded_file = st.file_uploader(
        "Upload your resume",
        type=["pdf", "docx", "txt"],
        help="Supported formats: PDF, DOCX, and TXT (maximum 10 MB).",
    )
    job_description = st.text_area(
        "Paste the job description (optional)",
        height=220,
        placeholder="Add a job description to see matching and missing keywords.",
    )

    if uploaded_file is None:
        st.info("Upload a resume to get started.")
        return

    if uploaded_file.size > MAX_FILE_SIZE:
        st.error("This file is larger than 10 MB. Please upload a smaller resume.")
        return

    try:
        resume_text = extract_resume_text(uploaded_file.name, uploaded_file.getvalue())
    except (
        BadZipFile,
        KeyError,
        OSError,
        PackageNotFoundError,
        PdfReadError,
        UnicodeDecodeError,
        ValueError,
    ) as error:
        st.error(f"Could not read this resume: {error}")
        return

    if not resume_text.strip():
        st.error(
            "No readable text was found. If this is a scanned PDF, run OCR on it "
            "and upload a searchable version."
        )
        return

    analysis = analyze_resume(resume_text, job_description)
    if len(resume_text.strip()) < 100:
        st.warning(
            "This file contains very little text. Check that the complete resume "
            "was uploaded and that the PDF is searchable."
        )

    st.subheader("Your results")
    score_col, completeness_col, length_col = st.columns(3)
    if "match_score" in analysis:
        score_col.metric("Job match score", f"{analysis['match_score']} / 100")
    else:
        score_col.metric("Job match score", "Add a job description")
    completeness_col.metric("Resume completeness", f"{analysis['completeness_score']} / 100")
    length_col.metric("Resume length", f"{len(resume_text.split())} words")

    if "match_score" in analysis:
        matched_col, missing_col = st.columns(2)
        matched_col.markdown("**Matching keywords**")
        matched = analysis["matched_keywords"]
        matched_col.write(", ".join(matched) if matched else "No matching terms found.")
        missing_col.markdown("**Potentially missing keywords**")
        missing = analysis["missing_keywords"]
        missing_col.write(", ".join(missing) if missing else "No major gaps found.")

    st.subheader("Resume checklist")
    checklist_cols = st.columns(3)
    for index, (section, present) in enumerate(analysis["sections"].items()):
        checklist_cols[index % len(checklist_cols)].write(
            f"{'✅' if present else '⚠️'} {section}"
        )
    st.write(
        f"{'✅' if analysis['action_verb_count'] else '⚠️'} "
        f"Action verbs found: {analysis['action_verb_count']}"
    )
    if analysis["has_metrics"]:
        st.write("✅ Includes measurable results")
    else:
        st.write("⚠️ Add measurable results where possible")

    st.subheader("Suggestions")
    for recommendation in build_recommendations(analysis):
        st.markdown(f"- {recommendation}")

    st.caption(
        "Keyword and completeness scores are heuristic text checks. "
        "They do not assess design, factual accuracy, or a recruiter's judgment."
    )


if __name__ == "__main__":
    main()
