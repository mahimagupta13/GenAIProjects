import unittest
from io import BytesIO

from docx import Document

from main import analyze_resume, extract_resume_text


class ResumeCheckerTests(unittest.TestCase):
    def test_keyword_matching_uses_word_boundaries(self) -> None:
        analysis = analyze_resume("JavaScript developer", "Java")

        self.assertEqual(analysis["matched_keywords"], [])
        self.assertIn("java", analysis["missing_keywords"])

    def test_analysis_reports_matching_and_missing_terms(self) -> None:
        resume = """
        Name name@example.com
        Summary
        Experience
        Developed Python services and improved AWS deployments by 25%.
        Education
        Skills
        Python AWS project management
        Projects
        """

        analysis = analyze_resume(
            resume,
            "Python AWS project management SQL",
        )

        self.assertIn("python", analysis["matched_keywords"])
        self.assertIn("aws", analysis["matched_keywords"])
        self.assertIn("sql", analysis["missing_keywords"])
        self.assertGreater(analysis["completeness_score"], 0)
        self.assertGreater(analysis["match_score"], 0)

    def test_extracts_plain_text(self) -> None:
        text = extract_resume_text("resume.txt", b"Resume text")

        self.assertEqual(text, "Resume text")

    def test_extracts_docx_paragraphs_and_tables(self) -> None:
        document = Document()
        document.add_paragraph("Resume paragraph")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Resume table"
        content = BytesIO()
        document.save(content)

        text = extract_resume_text("resume.docx", content.getvalue())

        self.assertIn("Resume paragraph", text)
        self.assertIn("Resume table", text)


if __name__ == "__main__":
    unittest.main()
